from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "market_entry_section_word.docx"

FONT = "Malgun Gothic"
INK = RGBColor(17, 17, 17)
MUTED = RGBColor(95, 95, 95)
GREEN = RGBColor(118, 185, 0)
LIGHT_GREEN = "E9F6D1"
LIGHT_GRAY = "F2F2F2"
LINE = "D7D7D7"


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_paragraph_format(p, before=0, after=5, line=1.08):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=LINE, size="4"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_width(cell, cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    width = tc_pr.find(qn("w:tcW"))
    if width is None:
        width = OxmlElement("w:tcW")
        tc_pr.append(width)
    width.set(qn("w:w"), str(int(cm * 567)))
    width.set(qn("w:type"), "dxa")


def clear_cell(cell):
    cell.text = ""
    for paragraph in cell.paragraphs:
        paragraph._p.getparent().remove(paragraph._p)
    cell._tc.append(OxmlElement("w:p"))


def add_hyperlink(paragraph, text, url, size=8.2):
    part = paragraph.part
    rid = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT)
    r_fonts.set(qn("w:hAnsi"), FONT)
    r_fonts.set(qn("w:eastAsia"), FONT)
    r_pr.append(r_fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "4F7F00")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    r_pr.append(sz)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_para(doc, text="", style=None, size=10, bold=False, color=INK, after=5, before=0):
    p = doc.add_paragraph(style=style)
    set_paragraph_format(p, before=before, after=after)
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph()
    set_paragraph_format(p, before=14 if level == 1 else 10, after=6)
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, size=17, bold=True, color=INK)
    elif level == 2:
        set_run_font(run, size=14, bold=True, color=INK)
    else:
        set_run_font(run, size=12, bold=True, color=INK)
    return p


def add_callout(doc, text, label=None):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade_cell(cell, LIGHT_GREEN)
    set_cell_border(cell, color="B7D98C", size="6")
    clear_cell(cell)
    p = cell.paragraphs[0]
    set_paragraph_format(p, after=0, line=1.15)
    if label:
        r = p.add_run(label)
        set_run_font(r, 10, bold=True, color=INK)
    r = p.add_run(text)
    set_run_font(r, 10, color=INK)
    doc.add_paragraph()
    return table


def add_source_paragraph(cell, sources):
    p = cell.add_paragraph()
    set_paragraph_format(p, before=2, after=0, line=1.0)
    r = p.add_run("출처: ")
    set_run_font(r, 8.2, color=MUTED)
    for idx, (name, url) in enumerate(sources):
        if idx:
            r = p.add_run(", ")
            set_run_font(r, 8.2, color=MUTED)
        add_hyperlink(p, name, url, size=8.2)


def add_text_cell(cell, text, size=9.2, bold=False):
    clear_cell(cell)
    p = cell.paragraphs[0]
    set_paragraph_format(p, after=0, line=1.08)
    for idx, part in enumerate(text.split("\n")):
        if idx:
            p.add_run().add_break()
        r = p.add_run(part)
        set_run_font(r, size=size, bold=bold, color=INK)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP


def add_table(doc, headers, rows, widths, font_size=9.0):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        set_cell_width(hdr_cells[i], widths[i])
        shade_cell(hdr_cells[i], LIGHT_GRAY)
        set_cell_border(hdr_cells[i])
        add_text_cell(hdr_cells[i], header, size=9.2, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, item in enumerate(row):
            set_cell_width(cells[i], widths[i])
            set_cell_border(cells[i])
            if isinstance(item, dict):
                add_text_cell(cells[i], item["text"], size=font_size)
                if item.get("sources"):
                    add_source_paragraph(cells[i], item["sources"])
            else:
                add_text_cell(cells[i], str(item), size=font_size)
    doc.add_paragraph()
    return table


SOURCES = {
    "worldbank": ("World Bank MS.MIL.XPND.CD", "https://data.worldbank.org/indicator/MS.MIL.XPND.CD"),
    "france": ("프랑스 국방부·SCORPION 공개자료", "https://www.defense.gouv.fr/"),
    "poland": ("폴란드 국방부 공개자료", "https://www.gov.pl/web/national-defence"),
    "patria": ("Patria 6x6 CAVS 공개자료", "https://www.patriagroup.com/"),
    "iveco": ("Iveco Defence Vehicles Guarani 공개자료", "https://www.ivecodefencevehicles.com/"),
    "edge": ("EDGE Group 공개자료", "https://edgegroup.ae/"),
    "iiss": ("IISS 공개 장비자료", "https://www.iiss.org/"),
    "sipri": ("SIPRI Military Expenditure Database", "https://www.sipri.org/databases/milex"),
    "nato": ("NATO 방위비·억제태세 공개자료", "https://www.nato.int/"),
    "eda": ("European Defence Agency Defence Data", "https://eda.europa.eu/"),
    "ec": ("European Commission Defence Industry", "https://defence-industry-space.ec.europa.eu/"),
    "rand": ("RAND 무인체계·우크라이나 전쟁 교훈", "https://www.rand.org/"),
    "csis": ("CSIS 무인체계·우크라이나 전쟁 교훈", "https://www.csis.org/"),
    "infodefensa": ("Infodefensa 중남미 장갑차 공개보도", "https://www.infodefensa.com/"),
}


REGIONS = [
    {
        "title": "유럽",
        "lead": "유럽은 러시아 위협과 우크라이나 전쟁 교훈으로 인해 차륜형장갑차 현대화, 전술드론 대량운용, 대드론 방호, 전술망 연동 투자가 동시에 확대되고 있다. 특히 동유럽과 북유럽 주요국은 최근 5년간 국방비가 연평균 9~29% 수준으로 증가하여 군집드론 적용 장갑차의 초기 진입 가능성이 높다.",
        "rows": [
            ("프랑스", "국방비는 2020~2024년 연평균 약 4.9% 증가했으며, SCORPION 사업으로 VBMR Griffon 6x6 약 1,872대, EBRC Jaguar 약 300대, VBMR-L Serval 약 978대를 전력화하고 있다. Griffon의 지휘·정찰 파생형 구조를 고려하면 후속 개량 단계에서 군집드론 저장·충전·임무통제 모듈을 별도 파생형으로 제기할 수 있다.", [SOURCES["worldbank"], SOURCES["france"]]),
            ("폴란드", "국방비는 2020~2024년 연평균 약 29.0% 증가했으며, Rosomak 장갑차 현대화, Borsuk IFV 최대 약 1,400대 프레임워크, WB Group Gladius 약 20억 PLN 및 FlyEye 무인기 공급이 공개되어 있다. 장갑차 현대화와 무인체계 획득이 동시에 진행되므로 기계화부대용 군집드론 모함형 장갑차 소요를 직접 제기할 수 있다.", [SOURCES["worldbank"], SOURCES["poland"]]),
            ("루마니아", "국방비는 2020~2024년 연평균 약 14.6% 증가했으며, Bayraktar TB2 18대 약 3.21억 달러 계약, Piranha V 8x8 운용 확대, 298대급 IFV 획득 프로그램이 공개되어 있다. 흑해·우크라이나 인접지역 감시정찰 수요가 높아 차륜형장갑차 기반 군집드론 ISR·통신중계 임무모듈 제안이 가능하다.", [SOURCES["worldbank"], SOURCES["nato"]]),
            ("핀란드", "국방비는 2020~2024년 연평균 약 15.9% 증가했으며, Patria 6x6 CAVS 프로그램에서 2023년 91대 주문과 2024년 옵션 70대 행사가 공개되어 있다. 혹한 환경에서 군집드론 배터리 관리, 다수 드론 동시운용, 통신중계 기능을 검증할 수 있는 북유럽형 실증시장으로 활용 가능하다.", [SOURCES["worldbank"], SOURCES["patria"]]),
        ],
    },
    {
        "title": "중남미",
        "lead": "중남미는 유럽처럼 국방비가 급증하는 시장은 아니지만, 노후 장갑차 교체, 국경·정글·산악 감시, 현지조립형 산업협력 수요가 존재한다. 따라서 대형 양산보다 소규모 실증, 기존 장갑차 개량, 현지 조립형 군집드론 임무모듈이 현실적인 진입 방식이다.",
        "rows": [
            ("브라질", "국방비는 2020~2024년 연평균 약 1.7%로 완만히 증가했으며, VBTP-MR Guarani 6x6 계열은 공개자료상 2,000대 이상 장기계획으로 언급된다. 자국 생산 차륜형 플랫폼의 지휘·정찰 파생형 확대 구조를 활용하면 군집드론 통제모듈을 별도 파생형으로 제안할 수 있다.", [SOURCES["worldbank"], SOURCES["iveco"]]),
            ("칠레", "국방비는 2020~2024년 연평균 약 -0.7%로 정체되어 대규모 신규 도입보다 기존 Mowag/Piranha 계열 및 노후 장갑차의 유지·교체 수요가 중심이다. 따라서 기존 장갑차에 소형 군집드론 발사·통제 장비를 얹는 저위험 개량형 소요로 접근하는 것이 적합하다.", [SOURCES["worldbank"], SOURCES["infodefensa"]]),
            ("페루", "국방비는 2020~2024년 연평균 약 -0.8%로 정체되어 있으나, FAME-STX-Hyundai Rotem 군용차량 조립협력과 차륜형·궤도형 차량 현지 생산 협력이 공개되어 있다. 현지 조립·정비 협력 기반을 활용하면 산악·국경 감시용 군집드론 운용모듈을 소규모 실증사업으로 제기할 수 있다.", [SOURCES["worldbank"]]),
        ],
    },
    {
        "title": "중동",
        "lead": "중동은 드론·미사일 위협, 에너지 인프라 방호, 국경감시, 사막기동 요구가 결합된 시장이다. 국방비 규모가 크고 무인체계·대드론 투자가 활발하여, 군집드론 운용장갑차를 고성능 임무모듈·대드론 연동 패키지로 제안할 수 있다.",
        "rows": [
            ("사우디아라비아", "국방비는 2020~2024년 연평균 약 5.6% 증가했으며, Baykar Akıncı 대형 UAV 수출계약과 Vision 2030 기반 방산 현지화가 공개되어 있다. 사막 국경감시와 중요시설 방호에서 다수 드론을 동시에 운용하고 대드론·전자전 장비와 연동하는 군집드론 장갑차 소요가 크다.", [SOURCES["worldbank"]]),
            ("UAE", "World Bank 2020~2024 군사비 공개값은 공백이나, Rabdan 8x8 최대 약 700대·약 6.61억 달러 규모 사업과 NIMR 4x4·6x6 계열, EDGE Group 무인체계 투자가 공개되어 있다. 완제품보다 현지 공동개발형 군집드론 임무SW, 자동발사·충전모듈, 차량통합 패키지로 접근하는 것이 적합하다.", [SOURCES["worldbank"], SOURCES["edge"]]),
            ("카타르", "World Bank 공개값은 2021~2022년만 확인되며 해당 기간 연평균 약 32.9% 증가했다. VAB 계열 장갑차 현대화, Boxer RCT30 10대 주문·초도 납품, Altay 전차 100대 주문이 공개되어 있어 기지방호·국경감시·대드론 연동형 군집드론 운용모듈의 단가와 후속지원 가치가 높다.", [SOURCES["worldbank"], SOURCES["iiss"]]),
        ],
    },
]


ENTRY_SUMMARY = [
    ("폴란드", "2020~2024년 국방비 CAGR 약 29.0%. Rosomak 현대화, Borsuk IFV 최대 약 1,400대, Gladius 약 20억 PLN, FlyEye 공급.", "기계화부대와 같이 움직이는 군집드론 저장·충전·동시발사·임무통제 차량 소요로 직접 연결 가능."),
    ("루마니아", "2020~2024년 약 14.6%. TB2 18대 약 3.21억 달러, Piranha V, 298대급 IFV 사업 공개.", "흑해·동부전선 감시정찰 수요가 높아 다수 소형드론 반복 운용 플랫폼 수요가 명확."),
    ("발트 3국", "에스토니아 19.0%, 라트비아 17.7%, 리투아니아 22.3%. Patria 6x6, ARMA 6x6, Boxer Vilkas 획득.", "소량·분산형 신속획득에 적합하며 영토방위군용 차량형 발사·통제모듈로 접근 가능."),
    ("핀란드", "2020~2024년 약 15.9%. Patria 6x6 CAVS 91대 주문 및 70대 옵션 행사.", "혹한 환경 배터리 관리, 통신중계, 분산작전 통제 검증시장으로 적합."),
    ("노르웨이", "2020~2024년 약 9.6%. Leopard 2A8 54대+옵션 18대, CV90 현대화.", "북극권·해안·산악 감시정찰용 고신뢰 군집드론 임무모듈 시장 가능."),
    ("호주", "2020~2024년 약 5.5%. LAND 400 Phase 3 Redback IFV 129대, Ghost Shark·MQ-28 등 무인체계 투자.", "인도태평양 장거리 ISR과 유무인복합작전 수요를 지상군 군집드론 모함차량으로 확장 가능."),
    ("사우디아라비아", "2020~2024년 약 5.6%. Akıncı UAV 수출계약, Vision 2030 방산 현지화, 국경방호 전력 운용.", "사막 국경감시와 중요시설 방호에서 군집드론과 대드론·전자전 장비 연동 수요가 큼."),
    ("UAE", "World Bank 2020~2024 공개값 공백. Rabdan 8x8 최대 약 700대·약 6.61억 달러, EDGE·NIMR 투자.", "현지 공동개발형 군집드론 임무SW, 차량형 발사·충전모듈, 통제 패키지가 유리."),
    ("인도", "2020~2024년 약 4.2%. DRDO/Tata WhAP, FRCV 약 1,770대 장기 프로그램, V-BAT 현지생산 협력.", "고산·사막 국경에서 다수 드론을 반복 전개하는 현지생산형 임무모듈 가능성이 큼."),
    ("필리핀", "2020~2024년 약 6.9%. Sabrah 경전차 약 1.72억 달러, Guarani 6x6, Hermes 900/450 UAS 약 1.75억 달러.", "도서·해안 감시와 회색지대 대응을 위한 소형 군집드론 신속전개형 차량 소요."),
    ("페루", "2020~2024년 약 -0.8%. FAME-STX-Hyundai Rotem 군용차량 조립협력 공개.", "국방비 성장률은 낮지만 현지조립·산악감시 실증사업으로 접근 가능."),
    ("카타르", "2021~2022년 단기 CAGR 약 32.9%. VAB 현대화, Boxer RCT30 10대, Altay 100대.", "고성능 군집드론 임무모듈, 대드론 연동, 기지방호 패키지 단가가 높은 시장."),
]


COUNTRIES = [
    {
        "name": "폴란드",
        "rows": [
            ("차륜형장갑차 기반 군집드론 운용모듈", "시범 1개 대대급, 양산 전환 시 30~80대급 임무모듈", "2027~2035년 단계적 획득 가능"),
            ("전방 정찰·타격 드론 패키지", "여단급 드론 운용세트 및 차량 탑재형 통제장비", "2026~2032년 우선 도입 가능"),
        ],
        "intro": "폴란드는 NATO 동부전선의 핵심 국가이며, 우크라이나 전쟁 이후 장갑차·포병·군집 정찰드론·배회형 타격체계·방공 분야 투자를 동시에 확대하고 있다.",
        "background": "대규모 지상전 위협에 직접 노출되어 있고, 기동부대 단위의 정찰·타격·대드론 기능 통합 필요성이 높다.",
        "geo": "러시아 및 벨라루스와 인접하고 우크라이나 지원 후방기지 역할을 수행하고 있어 전방 감시정찰과 분산기동 생존성이 핵심 과제이다.",
        "concept": "차륜형장갑차 기반 군집드론 운용차량은 기계화부대와 함께 이동하면서 다수 드론의 동시발사, ISR, 표적획득, 통신중계, 포병 좌표전송을 수행하는 임무에 적합하다.",
        "conclusion": "위협수준, 예산 증가, 장갑차 현대화, 군집드론 운용 확대 가능성이 모두 확인되므로 군집드론 적용 가능 장갑차 개발소요를 제기하기 가장 적합한 1순위 진입 추진국이다.",
    },
    {
        "name": "루마니아",
        "rows": [
            ("국경·흑해 연안 군집드론 운용장갑차", "초도 10~30대급, 국경감시부대 중심 확대 가능", "2027~2034년 단계적 도입 가능"),
            ("차륜형 장갑차 군집드론 임무모듈 개량", "전술통신·소형드론·대드론 센서 통합 패키지", "2026~2032년 우선소요 가능"),
        ],
        "intro": "루마니아는 흑해와 우크라이나에 인접한 NATO 동부전선 국가로, 감시정찰과 기동방호 수요가 높다.",
        "background": "동부 국경과 흑해 연안 방호를 위해 장갑차, 전술드론, 지휘통신체계를 묶은 패키지형 획득 가능성이 있다.",
        "geo": "흑해 지역은 러시아 해·공중 위협과 우크라이나 전쟁 영향이 겹치는 공간으로, 지속 감시와 신속 대응 능력이 요구된다.",
        "concept": "차량형 군집드론 운용 플랫폼은 국경감시, 해안 접근로 정찰, 다수 소형드론 동시 운용, 전방부대 통신중계, 포병 표적획득에 활용될 수 있다.",
        "conclusion": "폴란드보다 시장 규모는 작지만 NATO 동부전선 감시정찰 수요가 명확하므로 군집드론 운용장갑차의 중기 진입 추진국으로 적합하다.",
    },
    {
        "name": "발트 3국",
        "rows": [
            ("경량 차륜형 군집드론 운용 플랫폼", "국가별 5~20대급 소규모 패키지", "2026~2032년 신속도입 가능"),
            ("영토방위군용 군집드론·대드론 패키지", "분산부대용 군집드론 세트 및 차량형 통제장비", "2026~2030년 우선소요 가능"),
        ],
        "intro": "에스토니아·라트비아·리투아니아는 러시아 접경의 소규모 NATO 국가로, 분산방어와 소형 무인체계 운용에 대한 관심이 높다.",
        "background": "대형 장갑차 대량도입보다 소형·모듈형·저비용 드론 운용 패키지의 현실성이 높다.",
        "geo": "짧은 전략종심과 러시아 접경이라는 조건 때문에 조기경보, 국경감시, 분산부대 생존성이 중요하다.",
        "concept": "차량형 군집드론 운용체계는 소규모 방어거점 간 이동, 전방 감시, 통신중계, 영토방위군 지원 임무에 적합하다.",
        "conclusion": "수량은 제한적이나 소형 군집드론의 분산운용 실증과 NATO 공동조달 연계 가능성이 있어 기술검증형 진입국으로 의미가 있다.",
    },
    {
        "name": "핀란드",
        "rows": [
            ("혹한형 차륜형장갑차 군집드론 운용모듈", "초도 10~30대급 시험운용, 후속 확대 가능", "2027~2035년 단계적 획득 가능"),
            ("분산작전용 군집 정찰·통신중계 드론 패키지", "여단급 ISR 및 통신중계 세트", "2026~2032년 도입 가능"),
        ],
        "intro": "핀란드는 러시아와 긴 국경을 맞대고 있으며, 혹한·산림·분산작전 환경에서 무인체계 활용 필요성이 높다.",
        "background": "NATO 가입 이후 동맹 차원의 방어태세 강화와 기존 지상전력 현대화 수요가 결합되어 있다.",
        "geo": "북유럽 접경방어 환경은 넓은 감시구역과 제한된 접근로가 특징이므로 무인 감시정찰과 통신중계가 중요하다.",
        "concept": "혹한형 차량 개조, 배터리 관리, 통신중계 드론, 분산부대 영상공유 기능을 강조해야 한다.",
        "conclusion": "혹한에서 군집드론 배터리·통신·동시운용 안정성을 검증하면 북유럽·NATO 시장으로 확장 가능한 전략적 진입국이다.",
    },
    {
        "name": "노르웨이",
        "rows": [
            ("북극권 감시정찰형 군집드론 운용차량", "소량 고성능 5~20대급 패키지", "2027~2035년 도입 가능"),
            ("해안·산악 방호용 군집드론 ISR 임무모듈", "전술통신, 군집드론 통제, 상황공유 장비 통합", "2026~2033년 개량 가능"),
        ],
        "intro": "노르웨이는 북극권과 해안 방어, NATO 북부전선 감시정찰 수요가 큰 국가이다.",
        "background": "대량획득보다는 고신뢰 장비, 혹한 운용, 네트워크 연동, 정밀 감시정찰 능력에 대한 수요가 높다.",
        "geo": "러시아 북방함대와 북극권 군사활동 증가는 해안·산악·극지 감시정찰의 중요성을 높이고 있다.",
        "concept": "차량형 군집드론 운용체계는 해안 접근로 감시, 산악지역 전방정찰, 지휘소와의 영상공유 임무에 적합하다.",
        "conclusion": "수량은 제한적이나 고신뢰 군집드론 임무모듈과 혹한형 차량통합 기술을 앞세워 고부가가치 시장으로 접근할 수 있다.",
    },
    {
        "name": "호주",
        "rows": [
            ("장거리 감시정찰용 군집드론 운용차량", "초도 10~40대급, 기갑·정찰부대 중심 확대 가능", "2027~2035년 도입 가능"),
            ("인도태평양 전개형 군집드론 임무모듈", "통신중계, 군집드론 통제, 열대·건조지역 운용 패키지", "2026~2034년 시험 가능"),
        ],
        "intro": "호주는 인도태평양 장거리 감시정찰과 분산전개 능력을 중시하며, 한국 방산과의 협력 경험도 보유하고 있다.",
        "background": "넓은 작전구역, 장거리 기동, 네트워크 중심전, 무인체계 통합 수요가 존재한다.",
        "geo": "인도태평양 안보경쟁과 북부지역 방어 수요는 지상부대의 감시정찰 범위 확대를 요구한다.",
        "concept": "군집드론 운용차량은 정찰부대의 전방 ISR, 다수 드론 동시운용, 통신중계, 원거리 표적획득, 다국적 연합훈련 지원에 활용될 수 있다.",
        "conclusion": "기존 협력 기반을 활용하면 유무인복합작전용 군집드론 운용장갑차의 중장기 공동개발 후보로 검토할 가치가 높다.",
    },
    {
        "name": "사우디아라비아",
        "rows": [
            ("사막형 군집드론 운용장갑차", "초도 20~60대급, 국경방호·기계화부대 확대 가능", "2027~2035년 단계적 획득 가능"),
            ("대드론·전자전 연동 임무모듈", "중요시설·국경·기지방호용 고부가 패키지", "2026~2032년 우선소요 가능"),
        ],
        "intro": "사우디아라비아는 드론·미사일 위협 경험과 높은 국방투자 여력을 보유한 중동 핵심 시장이다.",
        "background": "국경방호, 에너지 시설 방호, 사막기동, 대드론 대응 수요가 동시에 존재한다.",
        "geo": "예멘, 홍해, 걸프 지역 위협은 소형드론·미사일·무인수상정 등 복합 위협 대응을 요구한다.",
        "concept": "차량형 플랫폼은 사막 국경 순찰, 다수 정찰드론 동시발진, 대드론 센서 연동, 기지 외곽 방호에 적합하다.",
        "conclusion": "현지화 요구가 높을 수 있으나 사막형 군집드론 운용장갑차와 대드론 연동 패키지의 사업규모가 커 전략시장으로 분류할 수 있다.",
    },
    {
        "name": "UAE",
        "rows": [
            ("AI 기반 군집드론 운용 임무모듈", "시범 5~20대급, 현지 공동개발 시 확대 가능", "2026~2033년 추진 가능"),
            ("수출형 군집드론 통합 플랫폼", "차량·드론·통제SW 공동개발 패키지", "2027~2035년 공동사업 가능"),
        ],
        "intro": "UAE는 무인체계, AI, 방산 현지화에 적극적인 국가로, 완제품 판매보다 공동개발·현지화 모델이 유리하다.",
        "background": "국방기술 투자와 현지 방산기업 육성 정책이 강해 군집드론 운용장갑차를 기술협력 패키지로 제안할 수 있다.",
        "geo": "걸프 지역 긴장, 해상교통로 방호, 중요 인프라 보호 수요가 지속된다.",
        "concept": "차량형 드론 플랫폼은 사막·해안 감시, 신속대응부대 지원, 대드론 체계와의 데이터 연동 임무에 적합하다.",
        "conclusion": "UAE는 수량보다 군집드론 임무SW, 차량형 자동발사·충전모듈, 제3국 수출 연계 가능성을 중심으로 접근해야 한다.",
    },
    {
        "name": "인도",
        "rows": [
            ("고산·사막형 차륜형 군집드론 운용 플랫폼", "시범 20~50대급, 양산 시 대규모 확대 가능", "2027~2036년 장기 획득 가능"),
            ("국경감시·표적획득 군집드론 패키지", "군단·사단급 군집 정찰드론 및 차량형 통제장비", "2026~2034년 단계적 도입 가능"),
        ],
        "intro": "인도는 중국·파키스탄과의 국경분쟁, 고산지·사막지형, 대규모 지상군 운용으로 인해 군집드론 운용장갑차 수요가 클 수 있다.",
        "background": "대규모 병력과 긴 국경선, 현지 생산 요구가 결합되어 장기 대형시장으로 평가된다.",
        "geo": "히말라야 고산지대와 서부 사막지역은 감시정찰 공백을 줄이고 부대 생존성을 높이는 무인체계가 필요하다.",
        "concept": "고산·사막 환경에서 차량이 드론을 전개하고, 통신중계와 표적획득을 수행하는 방식이 적합하다.",
        "conclusion": "기술이전과 현지생산 요구가 크지만, 고산·사막 국경용 군집드론 운용장갑차로 채택될 경우 가장 큰 장기시장 중 하나가 될 수 있다.",
    },
    {
        "name": "필리핀",
        "rows": [
            ("도서·해안 감시정찰형 군집드론 운용차량", "초도 5~20대급, 해병·육군 신속대응부대 중심", "2027~2034년 도입 가능"),
            ("경량 군집드론 통제 패키지", "도서방어·내부안보 작전용 소형 패키지", "2026~2032년 우선소요 가능"),
        ],
        "intro": "필리핀은 도서·해안 감시, 회색지대 대응, 반군작전 수요가 결합되어 경량형 군집드론 운용 플랫폼에 적합하다.",
        "background": "대형 장갑차보다 신속전개형 차량과 소형 군집드론 통제장비의 실용성이 높다.",
        "geo": "남중국해 긴장과 다수 도서 방어 환경은 감시정찰 공백을 줄이는 이동형 무인체계 운용을 요구한다.",
        "concept": "차량형 플랫폼은 해안 접근로 감시, 도서 내 신속대응, 소형드론 전개, 지휘소 영상공유에 적합하다.",
        "conclusion": "시장규모는 중간 수준이나 도서·해안 감시를 위한 소형 군집드론 운용개념 적합성이 높아 경량형 파생모델 진입국으로 검토할 수 있다.",
    },
    {
        "name": "페루",
        "rows": [
            ("산악·국경 감시정찰형 군집드론 운용차량", "초도 5~20대급, 국경감시부대 중심 확대 가능", "2027~2035년 단계적 추진 가능"),
            ("노후 장갑차 개량 연계 군집드론 임무모듈", "정찰·통신중계·군집드론 통제 패키지", "2026~2033년 개량 가능"),
        ],
        "intro": "페루는 산악·정글·국경 지역 감시수요와 노후 지상장비 현대화 수요가 결합되어 있다.",
        "background": "중남미에서 대규모 전면전 가능성은 낮지만, 넓은 국경과 지형 제한 때문에 소형 드론과 차량형 통제장비의 실효성이 높다.",
        "geo": "안데스 산악지형과 정글지역은 지상 감시가 어렵고, 국경·치안·군 합동작전 수요가 지속된다.",
        "concept": "차량형 군집드론 운용체계는 산악도로와 전방기지에서 다수 정찰드론을 전개하고, 지휘소와 영상을 공유하는 임무에 적합하다.",
        "conclusion": "페루는 대형 양산시장보다는 산악·국경 감시용 군집드론 운용모듈을 현지조립·개량형 패키지로 실증하는 진입국으로 검토하는 것이 타당하다.",
    },
    {
        "name": "카타르",
        "rows": [
            ("사막형 차륜형장갑차 군집드론 운용모듈", "초도 10~20대급, 주요 기지·국경방호 부대 중심 확대 가능", "2027~2032년 도입 가능"),
            ("군집드론·대드론·감시정찰 통합 패키지", "공군기지, 에너지 인프라, 국경감시용 차량 탑재형 패키지", "2026~2030년 우선소요 가능"),
        ],
        "intro": "카타르는 소규모 병력으로 넓은 전략시설과 국경을 방호해야 하는 국가로, 고성능 장비와 임무자동화 수요가 높다.",
        "background": "방산 구매력이 높고 기존 장갑차 현대화와 대드론·기지방호·국경감시 수요가 명확하다.",
        "geo": "걸프 지역은 드론, 순항미사일, 탄도미사일, 해상위협이 복합적으로 존재한다.",
        "concept": "차량 탑재형 군집드론 운용모듈은 사막 국경감시, 기지 외곽 정찰, 중요시설 방호, 대드론 탐지체계 연동에 적합하다.",
        "conclusion": "수량은 제한적일 수 있으나 고성능 군집드론 운용모듈, 대드론 연동, 기지방호 패키지의 단가와 후속지원 가치가 높은 중동 우선 진입국이다.",
    },
]


def add_labeled_paragraph(doc, label, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, before=1, after=4, line=1.08)
    r = p.add_run(f"{label}: ")
    set_run_font(r, 9.8, bold=True, color=INK)
    r = p.add_run(text)
    set_run_font(r, 9.8, color=INK)


def build():
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.65)
    section.right_margin = Cm(1.65)
    section.top_margin = Cm(1.75)
    section.bottom_margin = Cm(1.55)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10)
    normal.font.color.rgb = INK

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(title, after=7)
    r = title.add_run("차륜형장갑차 기반 군집드론 운용 플랫폼 해외시장 분석")
    set_run_font(r, 20, bold=True, color=INK)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(subtitle, after=12)
    r = subtitle.add_run("2. 개조개발 필요성 작성용 Word 문서본")
    set_run_font(r, 11, bold=True, color=MUTED)

    add_callout(
        doc,
        "본 문서는 HTML 분리본의 내용을 Word/HWP 편집에 맞게 재배치한 문서이다. 특정 차종 적용 표현은 제외하고, 공개자료 기반의 시장성·획득사업·진입국가 분석만 정리하였다.",
    )

    add_heading(doc, "2. 개조개발 필요성", 1)
    add_para(doc, "러시아-우크라이나 전쟁 이후 지상전에서는 소형 드론, FPV 드론, 배회형 탄약, 대드론 장비가 동시에 운용되고 있다. 이에 따라 장갑차는 단순 병력수송 수단을 넘어, 전방에서 다수 드론을 저장·충전·발사·통제하고 상급 지휘체계와 데이터를 연동하는 임무 플랫폼으로 확장될 필요가 있다.")
    add_para(doc, "기존 차륜형장갑차 기반 개조개발은 신규 차체 개발보다 기간과 위험을 줄일 수 있으며, 후방 병력공간을 임무공간으로 전환하여 드론 저장고, 자동발사기, 충전장치, 데이터링크, 임무컴퓨터를 통합할 수 있다. 특히 중형 차륜형 플랫폼은 기동부대와 같은 속도로 이동하면서 전방 ISR, 통신중계, 표적획득, 제한적 타격 임무를 수행할 수 있어 유무인복합작전의 현실적인 출발점이 된다.")

    add_heading(doc, "2.1 해외시장 분석", 1)
    add_callout(doc, "향후 10년간 세계 지상무기 시장은 장갑차 신규 도입보다 기존 차륜형장갑차의 임무모듈화, 드론 운용능력 추가, 대드론 방호 통합, 네트워크 중심전 연동을 중심으로 성장할 가능성이 높다.")

    add_table(
        doc,
        ["구분", "주요 내용"],
        [
            ("수요 요인", "드론 위협 증가, 전방 감시정찰 확대, 병력 손실 최소화, 분산작전 확대"),
            ("획득 방식", "신규 플랫폼 도입, 기존 장갑차 개조, 임무모듈 구매, 드론·통신장비 패키지 통합"),
            ("우선 시장", "러시아 위협권 유럽, 사막·해안 방호 중동, 노후 장갑차 교체 수요가 있는 중남미"),
        ],
        [3.0, 14.7],
    )

    add_heading(doc, "2.1.1 세계시장전망 및 예상되는 주요 획득사업", 1)
    for region in REGIONS:
        add_heading(doc, region["title"], 2)
        add_para(doc, region["lead"])
        rows = []
        for country, text, sources in region["rows"]:
            rows.append((country, {"text": text, "sources": sources}))
        add_table(doc, ["주요 획득국", "주요 획득사업(예상)"], rows, [3.0, 14.7], font_size=8.8)

    add_heading(doc, "2.1.1.1 진입 추진국가 선정 및 분석", 1)
    add_callout(
        doc,
        "진입 추진국가는 위협수준, 군집드론 운용 가능성, 차륜형 플랫폼 보유 여부, 향후 도입 가능성, 시장성을 기준으로 선정하였다. 예비분석 후보 전체를 포함하되, 각국의 실제 사업명은 공개자료로 확정되지 않은 경우 ‘예상’으로 표기하였다.",
    )
    add_callout(
        doc,
        "주요 진입 후보국은 이미 장갑차 현대화, 전술드론 도입, 대드론 방호, 전술통신 고도화를 병행하고 있다. 따라서 본 사업의 수출·협력 논리는 ‘신규 장갑차 판매’보다 ‘기존 또는 신규 차륜형장갑차에 군집드론 저장·충전·자동발사·임무통제 기능을 통합하는 임무모듈 개발소요’로 제기하는 것이 타당하다.",
        label="개발소요 제기 핵심 결론: ",
    )

    add_heading(doc, "진입 추진국가 요약", 2)
    add_para(
        doc,
        "주: CAGR은 World Bank 군사비 지출 지표(MS.MIL.XPND.CD)의 명목 미 달러 기준 최근 공개값으로 계산하였다. 군집드론 적용 장갑차 시장은 직접 통계가 제한되므로 국방비 성장률, 실제 장갑차 획득사업, 무인체계 획득사업을 결합해 개발소요 가능성을 판단하였다.",
        size=9,
        color=MUTED,
    )
    add_table(
        doc,
        ["국가", "공개 확인 근거", "군집드론 장갑차 소요 연결"],
        ENTRY_SUMMARY,
        [2.55, 7.55, 7.6],
        font_size=8.1,
    )

    for country in COUNTRIES:
        add_heading(doc, country["name"], 2)
        add_table(
            doc,
            ["장비명", "사업 규모(예상)", "사업 일정(예상)"],
            country["rows"],
            [6.2, 6.0, 5.5],
            font_size=8.8,
        )
        add_para(doc, country["intro"], size=9.8)
        add_labeled_paragraph(doc, "선정배경", country["background"])
        add_labeled_paragraph(doc, "지정학적 상황 분석", country["geo"])
        add_labeled_paragraph(doc, "운용개념 분석", country["concept"])
        add_labeled_paragraph(doc, "결론", country["conclusion"])

    add_heading(doc, "출처 및 작성 기준", 1)
    add_para(doc, "본 분석은 공개자료를 바탕으로 한 기획단계 예측이며, 각국의 실제 획득사업은 예산, 정권, 안보상황, 현지 생산 요구, 수출통제에 따라 변경될 수 있다.")
    source_rows = [
        ("European Commission", "European Defence Industrial Strategy / EDIP", SOURCES["ec"][1]),
        ("European Defence Agency", "Defence Data Portal", SOURCES["eda"][1]),
        ("NATO", "Defence expenditure and deterrence posture", SOURCES["nato"][1]),
        ("World Bank", "Military expenditure current US$ indicator", SOURCES["worldbank"][1]),
        ("SIPRI", "Military Expenditure Database", SOURCES["sipri"][1]),
        ("French Ministry of Armed Forces", "SCORPION / Griffon, Jaguar, Serval 공개자료", SOURCES["france"][1]),
        ("Polish Ministry of National Defence", "Polish Armed Forces modernization 공개자료", SOURCES["poland"][1]),
        ("Patria", "Patria 6x6 / CAVS programme", SOURCES["patria"][1]),
        ("Iveco Defence Vehicles", "VBTP-MR Guarani 공개자료", SOURCES["iveco"][1]),
        ("EDGE Group", "EDGE / NIMR / Rabdan 공개자료", SOURCES["edge"][1]),
        ("IISS", "Qatar Armed Forces 장비 공개자료", SOURCES["iiss"][1]),
        ("Infodefensa", "칠레·중남미 장갑차 교체 공개보도", SOURCES["infodefensa"][1]),
        ("RAND / CSIS", "무인체계 및 우크라이나 전쟁 교훈", "https://www.rand.org/"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, h in enumerate(["기관·자료", "활용 내용", "바로가기"]):
        cell = table.rows[0].cells[i]
        set_cell_width(cell, [4.8, 7.2, 5.7][i])
        shade_cell(cell, LIGHT_GRAY)
        set_cell_border(cell)
        add_text_cell(cell, h, size=9.2, bold=True)
    for name, usage, url in source_rows:
        cells = table.add_row().cells
        widths = [4.8, 7.2, 5.7]
        for i, cell in enumerate(cells):
            set_cell_width(cell, widths[i])
            set_cell_border(cell)
        add_text_cell(cells[0], name, size=8.5)
        add_text_cell(cells[1], usage, size=8.5)
        clear_cell(cells[2])
        p = cells[2].paragraphs[0]
        set_paragraph_format(p, after=0, line=1.0)
        add_hyperlink(p, url.replace("https://", ""), url, size=8.3)

    add_para(doc, "표의 사업 규모와 일정은 ‘예상’이며, 실제 제안서 반영 시에는 각국 국방부 예산문서, 입찰공고, DSCA 또는 현지 조달기관 자료로 재검증해야 한다.", size=9, color=MUTED, before=8)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(footer, after=0)
    r = footer.add_run("공개자료 기반 작성본 | 2026.06.10")
    set_run_font(r, 8, color=MUTED)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
